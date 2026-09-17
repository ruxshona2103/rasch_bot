"""🆕 Admin uchun test natijalarini Excel (.xlsx) formatida eksport qilish."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

_HEADERS = [
    "O'rin",
    "ID",
    "F.I.Sh.",
    "Telegram username",
    "Telegram ID",
    "Telefon",
    "Viloyat",
    "Test turi",
    "Holati",
    "Ball (/75)",
    "Daraja",
    "Boshlagan vaqti",
    "Tugatgan vaqti",
]

_STATUS_LABELS = {
    "davom_etmoqda": "Davom etmoqda",
    "yakunlangan": "Yakunlangan",
    "vaqt_tugagan": "Vaqt tugagan",
}


def build_results_excel(test_title: str, rows: list[tuple]) -> bytes:
    """rows: [(attempt, user), ...] — db.queries.list_attempts_with_users_for_export
    natijasi. Qaytaradi: .xlsx fayl baytlari."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"

    ws.append([f"\"{test_title}\" — test natijalari"])
    ws.append([f"Jami ishtirokchilar: {len(rows)} ta"])
    ws.append([])
    ws.append(_HEADERS)

    header_row = 4
    for col in range(1, len(_HEADERS) + 1):
        ws.cell(row=header_row, column=col).font = Font(bold=True)

    for attempt, user in rows:
        ws.append(
            [
                attempt.rank_position or "",
                user.public_id or "",
                user.full_name,
                f"@{user.username}" if user.username else "",
                user.telegram_id,
                user.phone or "",
                user.region or "",
                "Jonli" if attempt.kind == "jonli" else "Arxiv",
                _STATUS_LABELS.get(attempt.status, attempt.status),
                attempt.ball_75 if attempt.ball_75 is not None else "",
                attempt.grade or "",
                attempt.started_at.strftime("%d.%m.%Y %H:%M") if attempt.started_at else "",
                attempt.finished_at.strftime("%d.%m.%Y %H:%M") if attempt.finished_at else "",
            ]
        )

    widths = [8, 8, 26, 18, 14, 15, 14, 9, 15, 11, 8, 17, 17]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
